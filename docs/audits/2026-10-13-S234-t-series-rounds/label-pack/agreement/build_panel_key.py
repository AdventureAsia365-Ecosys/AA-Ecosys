#!/usr/bin/env python3
"""AA-783 — build the 4-lens panel key for the 40 S234 label-pack pieces (one-off, offline, no DB, no LLM).

Inputs (S234 scratch, gitignored, read once):
  .tmp-session/s234/panel/synth.py               round 1: per-lens verdicts, overall, fidelity, ai_tells, audit
  .tmp-session/s234/panel/synth_r2_data.py       round 2: same, list form
  .tmp-session/s234/panel/r3/synth_r3.py         round 3: idx -> (verdict, overall, fidelity, ai_tells)
  .tmp-session/s234/panel/r4/synth_r4.py         round 4: same; audit from panel/fidelity_r4.json
  .tmp-session/s234/final/joined.json            all 80 pieces: stored majority verdict, votes, deciding defects
  label-pack/S234-label-40.xlsx, sheet `_key`    the 40 piece ids -> round, index in round, gate status

Only top-level literal assignments are read from the synth scripts (ast.literal_eval); nothing is executed.

Output: panel_key.json (next to this script), then `--embed` rewrites the PANEL_KEY block of aa783_export.py.

    python3 build_panel_key.py [--scratch /home/nghiep/projects/AA-Ecosys/.tmp-session/s234] [--embed]
"""
from __future__ import annotations

import argparse
import ast
import hashlib
import json
import statistics as st
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
PACK = HERE.parent / "S234-label-40.xlsx"
LENSES = ("E", "B", "C", "K")  # editor, brand lead, channel specialist, paying client
LENS_NAMES = {"E": "editor", "B": "brand_lead", "C": "channel", "K": "client"}
SEVERITY = ("as_is", "light", "heavy", "reject")
LETTER = {"A": "as_is", "L": "light", "H": "heavy", "R": "reject"}


def _norm(v: str) -> str:
    v = v.strip()
    if v in LETTER:
        return LETTER[v]
    v = v.lower().split()[0]
    if v not in SEVERITY:
        raise ValueError(f"unknown verdict {v!r}")
    return v


def majority(votes: list[str]) -> str:
    """S234 rule: as-is needs >=3 as-is votes, light-or-better needs >=3 of 4, a 2-2 tie takes the
    stricter label. For 4 votes that is the upper median by severity."""
    return sorted(votes, key=SEVERITY.index)[2]


def literals(path: Path) -> dict:
    out = {}
    for node in ast.parse(path.read_text()).body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
            v = node.value
            # round 2 writes `E = "light light …".split()`: a str constant's .split() is still a literal
            if (isinstance(v, ast.Call) and isinstance(v.func, ast.Attribute) and v.func.attr == "split"
                    and not v.args and isinstance(v.func.value, ast.Constant) and isinstance(v.func.value.value, str)):
                out[node.targets[0].id] = v.func.value.value.split()
                continue
            try:
                out[node.targets[0].id] = ast.literal_eval(v)
            except ValueError:
                pass  # computed values (json.load, comprehensions) are skipped
    return out


def rounds(scratch: Path) -> dict[int, dict[int, dict]]:
    """round -> idx -> {lens: {verdict, overall, fidelity, ai_tells}, audit}"""
    res: dict[int, dict[int, dict]] = {}
    r1 = literals(scratch / "panel/synth.py")
    res[1] = {i: {"lens": {l: {"verdict": _norm(r1[l][i]), "overall": r1["ov"][l][i],
                               "fidelity": r1["fid"][l][i], "ai_tells": r1["ai"][l][i]} for l in LENSES},
                  "audit": r1["audit"][i]} for i in range(20)}
    r2 = literals(scratch / "panel/synth_r2_data.py")
    pre = {"E": "e", "B": "b", "C": "c", "K": "k"}
    res[2] = {i: {"lens": {l: {"verdict": _norm(r2[l][i]), "overall": r2[pre[l] + "o"][i],
                               "fidelity": r2[pre[l] + "f"][i], "ai_tells": r2[pre[l] + "a"][i]} for l in LENSES},
                  "audit": r2["aud"][i]} for i in range(20)}
    r3 = literals(scratch / "panel/r3/synth_r3.py")
    r4 = literals(scratch / "panel/r4/synth_r4.py")
    r4["AUD"] = {x["idx"]: x["fidelity"] for x in json.loads((scratch / "panel/fidelity_r4.json").read_text())}
    for rn, d in ((3, r3), (4, r4)):
        src = {"E": d["E_"], "B": d["B"], "C": d["C"], "K": d["K"]}
        res[rn] = {i: {"lens": {l: dict(zip(("verdict", "overall", "fidelity", "ai_tells"),
                                            (_norm(src[l][i][0]), *src[l][i][1:]))) for l in LENSES},
                       "audit": d["AUD"][i]} for i in range(20)}
    return res


def label_pack_key() -> list[dict]:
    import openpyxl  # local only
    wb = openpyxl.load_workbook(PACK, read_only=True)
    ws = [w for w in wb.worksheets if w.title.startswith("_key")][0]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    label = {r[1]: r for r in wb["Label"].iter_rows(min_row=2, values_only=True)}
    out = []
    for r in rows:
        if not r[1]:
            continue
        lr = label[r[1]]
        out.append({"pack_no": int(r[0]), "piece_id": r[1], "round": int(str(r[2]).lstrip("r")),
                    "idx_in_round": int(r[3]), "gate_status": r[4], "gate_failed": r[5],
                    "tenant": lr[2], "channel": lr[3]})
    return out


def build(scratch: Path) -> dict:
    data = rounds(scratch)
    joined = {(j["r"], j["i"]): j for j in json.loads((scratch / "final/joined.json").read_text())}
    pieces, mismatches = {}, []
    for k in label_pack_key():
        rn, i = k["round"], k["idx_in_round"]
        d, j = data[rn][i], joined[(rn, i)]
        if j["pid"] != k["piece_id"]:
            raise SystemExit(f"piece id mismatch r{rn}#{i}: pack {k['piece_id']} vs joined {j['pid']}")
        votes = [d["lens"][l]["verdict"] for l in LENSES]
        letters = "".join({"as_is": "A", "light": "L", "heavy": "H", "reject": "R"}[v] for v in votes)
        if letters != j["votes"]:
            mismatches.append(f"r{rn}#{i}: synth votes {letters} vs joined {j['votes']}")
        stored = _norm(j["verdict"])
        computed = majority(votes)
        if stored != computed:
            mismatches.append(f"r{rn}#{i}: stored verdict {stored} vs computed {computed} (stored kept)")
        light_votes = sum(v in ("as_is", "light") for v in votes)
        pieces[k["piece_id"]] = {
            **k,
            "lenses": {LENS_NAMES[l]: d["lens"][l] for l in LENSES},
            "votes": letters,
            "majority": stored,
            "tie": light_votes == 2,
            "mean": {m: round(st.mean(d["lens"][l][m] for l in LENSES), 3)
                     for m in ("overall", "fidelity", "ai_tells")},
            "fidelity_audit": d["audit"],
            "deciding_defects": j.get("defects") or "",
            # the text the panel scored (S234 pull); the export flags a piece whose DB text has changed since
            "panel_text_sha256": hashlib.sha256((j.get("text") or "").encode("utf-8")).hexdigest(),
        }
    return {
        "schema": "aa783-panel-key/1",
        "source": "S234 T9 commercial panel rounds 1-4 (4 independent Claude Code lens agents + fidelity auditor)",
        "lenses": LENS_NAMES,
        "dimensions_per_lens": ["verdict", "overall", "fidelity", "ai_tells"],
        "dimensions_not_kept": ["brand_voice", "channel_fit", "hook", "clarity", "commercial_value"],
        "majority_rule": "upper median by severity: as_is needs >=3 as_is; light-or-better needs >=3 of 4; 2-2 takes the stricter",
        "scale": "scores /10, higher is better (ai_tells 10 = no AI tells)",
        "notes": mismatches,
        "pieces": pieces,
    }


EMBED_BEGIN = "# --- PANEL_KEY BEGIN (generated by build_panel_key.py --embed; do not edit) ---"
EMBED_END = "# --- PANEL_KEY END ---"


def embed(key: dict, script: Path) -> None:
    src = script.read_text()
    a, b = src.index(EMBED_BEGIN), src.index(EMBED_END)
    blob = json.dumps(key, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    block = f"{EMBED_BEGIN}\nPANEL_KEY_JSON = {blob!r}\n"
    script.write_text(src[:a] + block + src[b:])


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--scratch", default="/home/nghiep/projects/AA-Ecosys/.tmp-session/s234")
    ap.add_argument("--embed", action="store_true", help="also rewrite the PANEL_KEY block of aa783_export.py")
    a = ap.parse_args()
    key = build(Path(a.scratch))
    (HERE / "panel_key.json").write_text(json.dumps(key, ensure_ascii=False, indent=1, sort_keys=True) + "\n")
    p = key["pieces"].values()
    print(f"pieces {len(key['pieces'])}; majority "
          f"{ {v: sum(x['majority'] == v for x in p) for v in SEVERITY} }; notes {len(key['notes'])}")
    for n in key["notes"]:
        print("  note:", n)
    if a.embed:
        embed(key, HERE / "aa783_export.py")
        print("embedded into aa783_export.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
